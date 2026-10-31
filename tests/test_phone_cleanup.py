from io import StringIO

import pytest
from django.core.management import CommandError, call_command
from openpyxl import load_workbook

from accounts.management.commands.phone_cleanup import COLUMNS, suggest
from accounts.models import User
from crm.models import Client, ClientPhoneNumber

GOOD = "+9647715952998"
STRAY_ZERO = "+96407715952998"
TOO_SHORT = "+964771595299"


def cmd(*args):
    out = StringIO()
    call_command("phone_cleanup", *args, stdout=out)
    return out.getvalue()


def lead_with(company, phone, normalized=None):
    client = Client.objects.create(name="Lead", company=company)
    row = ClientPhoneNumber.objects.create(client=client, phone_number=GOOD[:-1] + str(client.pk % 10), phone_type="mobile", is_primary=True)
    ClientPhoneNumber.objects.filter(pk=row.pk).update(
        phone_number=phone, phone_normalized=normalized if normalized is not None else f"key{row.pk}"
    )
    Client.objects.filter(pk=client.pk).update(phone_number=phone)
    return client, row


def decide(path, decisions):
    book = load_workbook(path)
    sheet = book.active
    for row in sheet.iter_rows(min_row=2):
        key = (row[0].value, row[1].value)
        if key in decisions:
            action, new_phone = decisions[key]
            row[COLUMNS.index("action")].value = action
            row[COLUMNS.index("new_phone")].value = new_phone
    book.save(path)


@pytest.mark.parametrize(
    "raw,expected",
    [(STRAY_ZERO, GOOD), ("+964964771595299", ""), (TOO_SHORT, ""), ("+9649647715952998", GOOD), ("+abc", "")],
)
def test_suggest_only_repairs_what_has_one_meaning(raw, expected):
    assert suggest(raw) == expected


@pytest.mark.django_db
def test_export_finds_each_problem_once_and_keeps_numbers_as_text(company, tmp_path):
    invalid, _ = lead_with(company, TOO_SHORT)
    stray, _ = lead_with(company, STRAY_ZERO)
    lead_with(company, GOOD)
    path = tmp_path / "phones.xlsx"

    cmd("export", str(path))

    sheet = load_workbook(path).active
    rows = {r[1]: dict(zip(COLUMNS, r)) for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {invalid.id, stray.id}
    assert rows[invalid.id]["problem"] == "invalid"
    assert rows[stray.id]["problem"] == "stray_zero"
    assert rows[stray.id]["suggested_phone"] == GOOD
    assert sheet["E2"].number_format == "@"
    assert isinstance(sheet["E2"].value, str)


@pytest.mark.django_db
def test_export_shows_which_lead_holds_the_same_number(company, tmp_path):
    keeper, _ = lead_with(company, GOOD, normalized="9647715952998")
    clash, _ = lead_with(company, "07715952998")
    path = tmp_path / "phones.xlsx"

    cmd("export", str(path))

    row = dict(zip(COLUMNS, next(load_workbook(path).active.iter_rows(min_row=2, values_only=True))))
    assert row["id"] == clash.id
    assert row["problem"] == "same_number"
    assert f"lead {keeper.id}" in row["clashes_with"]


@pytest.mark.django_db
def test_apply_previews_until_commit_and_restores(company, tmp_path):
    lead, row = lead_with(company, STRAY_ZERO)
    path = tmp_path / "phones.xlsx"
    undo = tmp_path / "undo.json"
    cmd("export", str(path))
    decide(path, {("lead", lead.id): ("use_suggested", None)})

    assert "1 rows would apply" in cmd("apply", str(path))
    row.refresh_from_db()
    assert row.phone_number == STRAY_ZERO

    cmd("apply", str(path), "--commit", "--undo-file", str(undo))
    row.refresh_from_db()
    lead.refresh_from_db()
    assert row.phone_number == lead.phone_number == GOOD
    assert row.phone_normalized == GOOD[1:]

    cmd("restore", str(undo))
    row.refresh_from_db()
    lead.refresh_from_db()
    assert row.phone_number == lead.phone_number == STRAY_ZERO


@pytest.mark.django_db
def test_commit_requires_undo_file(tmp_path):
    with pytest.raises(CommandError):
        cmd("apply", str(tmp_path / "x.xlsx"), "--commit")


@pytest.mark.django_db
def test_blank_rows_are_left_alone(company, tmp_path):
    _, row = lead_with(company, TOO_SHORT)
    path = tmp_path / "phones.xlsx"
    cmd("export", str(path))

    cmd("apply", str(path), "--commit", "--undo-file", str(tmp_path / "undo.json"))

    row.refresh_from_db()
    assert row.phone_number == TOO_SHORT


@pytest.mark.django_db
def test_set_new_rejects_an_invalid_number(company, tmp_path):
    lead, row = lead_with(company, TOO_SHORT)
    path = tmp_path / "phones.xlsx"
    cmd("export", str(path))
    decide(path, {("lead", lead.id): ("set_new", "12345")})

    output = cmd("apply", str(path), "--commit", "--undo-file", str(tmp_path / "undo.json"))

    row.refresh_from_db()
    assert row.phone_number == TOO_SHORT
    assert "not valid" in output


@pytest.mark.django_db
def test_set_new_rejects_a_number_on_another_lead(company, tmp_path):
    lead_with(company, "+9647701234567", normalized="9647701234567")
    lead, row = lead_with(company, TOO_SHORT)
    path = tmp_path / "phones.xlsx"
    cmd("export", str(path))
    decide(path, {("lead", lead.id): ("set_new", "+9647701234567")})

    output = cmd("apply", str(path), "--commit", "--undo-file", str(tmp_path / "undo.json"))

    row.refresh_from_db()
    assert row.phone_number == TOO_SHORT
    assert "already on lead" in output


@pytest.mark.django_db
def test_remove_keeps_the_leads_other_number_and_restore_recreates_the_row(company, tmp_path):
    lead, row = lead_with(company, TOO_SHORT)
    other = ClientPhoneNumber.objects.create(client=lead, phone_number="+9647701234567", phone_type="work")
    path = tmp_path / "phones.xlsx"
    undo = tmp_path / "undo.json"
    cmd("export", str(path))
    decide(path, {("lead", lead.id): ("remove", None)})

    cmd("apply", str(path), "--commit", "--undo-file", str(undo))

    lead.refresh_from_db()
    other.refresh_from_db()
    assert not ClientPhoneNumber.objects.filter(pk=row.pk).exists()
    assert lead.phone_number == other.phone_number
    assert other.is_primary

    cmd("restore", str(undo))
    lead.refresh_from_db()
    assert ClientPhoneNumber.objects.get(pk=row.pk).phone_number == TOO_SHORT
    assert lead.phone_number == TOO_SHORT


@pytest.mark.django_db
def test_user_rows_set_and_clear(company, tmp_path):
    first = User.objects.create_user(username="u1", email="u1@test.com", password="testpass123", company=company)
    second = User.objects.create_user(username="u2", email="u2@test.com", password="testpass123", company=company)
    User.objects.filter(pk=first.pk).update(phone="99911")
    User.objects.filter(pk=second.pk).update(phone=TOO_SHORT)
    path = tmp_path / "phones.xlsx"
    cmd("export", str(path))
    assert load_workbook(path).active.max_row == 2  # "99911" is not dialable, so it is left out
    decide(path, {("user", second.id): ("set_new", GOOD)})

    cmd("apply", str(path), "--commit", "--undo-file", str(tmp_path / "undo.json"))

    second.refresh_from_db()
    assert second.phone == GOOD


@pytest.mark.django_db
def test_row_changed_since_export_is_skipped(company, tmp_path):
    lead, _ = lead_with(company, TOO_SHORT)
    path = tmp_path / "phones.xlsx"
    cmd("export", str(path))
    decide(path, {("lead", lead.id): ("remove", None)})
    ClientPhoneNumber.objects.filter(client=lead).update(phone_number=GOOD, phone_normalized=GOOD[1:])

    output = cmd("apply", str(path), "--commit", "--undo-file", str(tmp_path / "undo.json"))

    assert "changed since export" in output
    assert ClientPhoneNumber.objects.filter(client=lead, phone_number=GOOD).exists()
