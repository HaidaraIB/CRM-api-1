"""
Spreadsheet workflow for phones only a person can fix: invalid numbers, numbers
with a stray 0 after the country code, and one number on two leads.

    python manage.py phone_cleanup export phone_cleanup.xlsx
    python manage.py phone_cleanup apply phone_cleanup.xlsx
    python manage.py phone_cleanup apply phone_cleanup.xlsx --commit --undo-file cleanup_undo.json
    python manage.py phone_cleanup apply phone_cleanup.xlsx --accept-stray-zero
    python manage.py phone_cleanup restore cleanup_undo.json

--accept-stray-zero treats every blank stray_zero row as use_suggested. Otherwise,
pick an action per row in the file and leave the rest blank:
    use_suggested  store suggested_phone
    set_new        store new_phone
    remove         clear the number from this user or lead

The file and the undo file hold real numbers. Delete both when done.
"""

import json

import phonenumbers
from django.core.management.base import BaseCommand, CommandError
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from accounts.models import User
from crm.models import Client, ClientPhoneNumber
from integrations.services.phone_match import canonical_phone_key
from validation.phones import is_dialable, to_e164, to_wire_digits

from .audit_phone_numbers import classify

COLUMNS = (
    "type", "id", "company", "name", "current_phone", "problem",
    "clashes_with", "suggested_phone", "action", "new_phone",
)
WIDTHS = (8, 9, 24, 28, 20, 16, 28, 20, 15, 20)
ACTIONS = ("use_suggested", "set_new", "remove")


def canonical(phone: str) -> str:
    """The valid +E.164 spelling of a number, or "" if it is not valid."""
    try:
        parsed = phonenumbers.parse(to_e164(phone))
    except phonenumbers.NumberParseException:
        return ""
    if not phonenumbers.is_valid_number(parsed):
        return ""
    return phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164)


def suggest(phone: str) -> str:
    """Repair typos that have one possible meaning; "" when a person has to decide."""
    fixed = canonical(phone)
    if fixed:
        return fixed
    try:
        parsed = phonenumbers.parse(to_e164(phone))
    except phonenumbers.NumberParseException:
        return ""
    code, national = str(parsed.country_code), str(parsed.national_number)
    return canonical(f"+{code}{national[len(code):]}") if national.startswith(code) else ""


def lead_rows_with(client_id: int, phone: str):
    wire = to_wire_digits(phone)
    rows = ClientPhoneNumber.objects.filter(client_id=client_id)
    return [row for row in rows if to_wire_digits(row.phone_number) == wire]


def clashes_for(row: ClientPhoneNumber):
    """Other leads already holding this number under its correct key."""
    wire = to_wire_digits(row.phone_number)
    return ClientPhoneNumber.objects.filter(company_id=row.company_id, phone_normalized=wire).exclude(
        client_id=row.client_id
    ).select_related("client")


def mark_shared_numbers(leads: list[dict]) -> None:
    """Rows that would become the same number as each other once corrected."""
    groups = {}
    for lead in leads:
        if lead["problem"] != "invalid":
            groups.setdefault((lead["company"], to_wire_digits(lead["current_phone"])), []).append(lead)
    for group in groups.values():
        if len({lead["id"] for lead in group}) < 2:
            continue
        for lead in group:
            others = [f"lead {o['id']}: {o['name']}" for o in group if o["id"] != lead["id"]]
            lead.update(problem="same_number", suggested_phone="", clashes_with=", ".join(others))


def problems():
    for user in User.objects.exclude(phone__isnull=True).exclude(phone="").select_related("company"):
        if is_dialable(user.phone) and (classify(user.phone) == "invalid" or canonical(user.phone) not in ("", user.phone)):
            yield {
                "type": "user", "id": user.id, "company": getattr(user.company, "name", ""),
                "name": user.username, "current_phone": user.phone,
                "problem": "invalid" if classify(user.phone) == "invalid" else "stray_zero",
                "clashes_with": "", "suggested_phone": suggest(user.phone),
            }

    seen, leads = set(), []
    rows = ClientPhoneNumber.objects.select_related("client__company").order_by("client_id", "id")
    for row in rows.iterator():
        value = row.phone_number
        if not is_dialable(value):
            continue
        kind = classify(value)
        if kind not in ("invalid", "convertible", "local_07"):
            continue
        key = (row.client_id, to_wire_digits(value))
        if key in seen:
            continue
        seen.add(key)
        clashes = list(clashes_for(row)) if kind != "invalid" else []
        problem = "invalid" if kind == "invalid" else "same_number" if clashes else "stray_zero"
        leads.append({
            "type": "lead", "id": row.client_id, "company": getattr(row.client.company, "name", ""),
            "name": row.client.name, "current_phone": value, "problem": problem,
            "clashes_with": ", ".join(f"lead {c.client_id}: {c.client.name}" for c in clashes),
            "suggested_phone": "" if problem == "same_number" else suggest(value),
        })
    mark_shared_numbers(leads)
    yield from leads


def write_workbook(path: str) -> int:
    book = Workbook()
    sheet = book.active
    sheet.title = "Phones"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    count = 0
    for problem in problems():
        count += 1
        sheet.append([problem.get(column, "") for column in COLUMNS])
        for column in ("current_phone", "suggested_phone", "new_phone"):
            cell = sheet.cell(row=sheet.max_row, column=COLUMNS.index(column) + 1)
            cell.number_format = "@"
            cell.data_type = "s"
    for index, width in enumerate(WIDTHS):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    dropdown = DataValidation(type="list", formula1=f'"{",".join(ACTIONS)}"', allow_blank=True)
    sheet.add_data_validation(dropdown)
    dropdown.add(f"I2:I{max(sheet.max_row, 2)}")
    book.save(path)
    return count


def read_workbook(path: str) -> list[dict]:
    sheet = load_workbook(path).active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if tuple(header[: len(COLUMNS)]) != COLUMNS:
        raise CommandError("This is not a phone_cleanup export: the columns do not match.")
    return [dict(zip(COLUMNS, ("" if v is None else str(v).strip() for v in row))) for row in rows]


def target_phone(row: dict) -> str:
    if row["action"] == "use_suggested":
        return canonical(row["suggested_phone"])
    return canonical(row["new_phone"])


def apply_user(row: dict, phone: str | None, undo: list) -> str | None:
    user = User.objects.filter(pk=row["id"], phone=row["current_phone"]).first()
    if user is None:
        return "number changed since export"
    if phone and User.objects.exclude(pk=user.pk).filter(phone=phone).exists():
        return "another user already has this number"
    undo.append({"source": "user", "pk": user.pk, "phone": user.phone})
    User.objects.filter(pk=user.pk).update(phone=phone)
    return None


def apply_lead(row: dict, phone: str | None, undo: list) -> str | None:
    client = Client.objects.filter(pk=row["id"]).first()
    rows = lead_rows_with(row["id"], row["current_phone"]) if client else []
    if not rows:
        return "number changed since export"
    if phone:
        taken = ClientPhoneNumber.objects.filter(
            company_id=rows[0].company_id, phone_normalized=to_wire_digits(phone)
        ).exclude(pk__in=[r.pk for r in rows])
        if taken.exists():
            return f"already on lead {taken.first().client_id}"

    undo.append({"source": "lead", "pk": client.pk, "phone": client.phone_number})
    matches_lead = to_wire_digits(client.phone_number) == to_wire_digits(row["current_phone"])
    for phone_row in rows:
        undo.append({
            "source": "lead_phone", "pk": phone_row.pk, "client_id": phone_row.client_id,
            "phone": phone_row.phone_number, "phone_type": phone_row.phone_type,
            "is_primary": phone_row.is_primary, "notes": phone_row.notes, "company_id": phone_row.company_id,
        })
    if phone:
        ClientPhoneNumber.objects.filter(pk__in=[r.pk for r in rows]).update(
            phone_number=phone, phone_normalized=canonical_phone_key(phone)
        )
        if matches_lead:
            Client.objects.filter(pk=client.pk).update(phone_number=phone)
        return None

    was_primary = any(r.is_primary for r in rows)
    ClientPhoneNumber.objects.filter(pk__in=[r.pk for r in rows]).delete()
    remaining = ClientPhoneNumber.objects.filter(client_id=client.pk).order_by("id").first()
    if remaining and was_primary:
        ClientPhoneNumber.objects.filter(pk=remaining.pk).update(is_primary=True)
    if matches_lead:
        Client.objects.filter(pk=client.pk).update(phone_number=remaining.phone_number if remaining else None)
    return None


class Command(BaseCommand):
    help = "Export phones that need a person, then apply the decisions. Previews unless --commit is given."

    def add_arguments(self, parser):
        parser.add_argument("step", choices=("export", "apply", "restore"))
        parser.add_argument("path")
        parser.add_argument("--commit", action="store_true", help="Write the changes.")
        parser.add_argument("--undo-file", help="Where --commit records old values. Required with --commit.")
        parser.add_argument(
            "--accept-stray-zero", action="store_true", help="Treat blank stray_zero rows as use_suggested."
        )

    def handle(self, *args, **options):
        if options["step"] == "export":
            self.stdout.write(f"{write_workbook(options['path'])} rows written to {options['path']}.")
        elif options["step"] == "apply":
            self._apply(options)
        else:
            self._restore(options["path"])

    def _apply(self, options):
        if options["commit"] and not options["undo_file"]:
            raise CommandError("--commit needs --undo-file.")
        undo, done, skipped = [], 0, 0
        claimed = {}
        for row in read_workbook(options["path"]):
            if not row["action"] and options["accept_stray_zero"] and row["problem"] == "stray_zero":
                row["action"] = "use_suggested"
            action = row["action"]
            if not action:
                continue
            label = f"{row['type']} {row['id']}"
            if action not in ACTIONS:
                reason = f"unknown action {action!r}"
            else:
                phone = None if action == "remove" else target_phone(row)
                reason = None if action == "remove" or phone else "the number to store is not valid"
                # Each row previews alone, so catch two rows in the file taking one number.
                holder = claimed.get((row["company"], phone)) if phone else None
                if holder and holder != label:
                    reason = f"{holder} in this file takes the same number"
            if reason is None:
                reason = self._run(row, phone, undo, options["commit"])
                if reason is None and phone:
                    claimed[(row["company"], phone)] = label
            if reason:
                skipped += 1
                self.stdout.write(f"    {label}: skipped, {reason}")
            else:
                done += 1
        if options["commit"]:
            with open(options["undo_file"], "w", encoding="utf8") as handle:
                json.dump(undo, handle)
            self.stdout.write(f"Applied {done} rows, skipped {skipped}. Undo file holds real numbers.")
        else:
            self.stdout.write(f"Dry run: {done} rows would apply, {skipped} would be skipped. Nothing was written.")

    def _run(self, row, phone, undo, commit) -> str | None:
        handler = apply_user if row["type"] == "user" else apply_lead
        try:
            with transaction.atomic():
                reason = handler(row, phone, undo)
                if not commit:
                    transaction.set_rollback(True)
                return reason
        except IntegrityError:
            return "conflicts with an existing number"

    def _restore(self, path):
        with open(path, encoding="utf8") as handle:
            undo = json.load(handle)
        for entry in reversed(undo):
            if entry["source"] == "user":
                User.objects.filter(pk=entry["pk"]).update(phone=entry["phone"])
            elif entry["source"] == "lead":
                Client.objects.filter(pk=entry["pk"]).update(phone_number=entry["phone"])
            else:
                ClientPhoneNumber.objects.update_or_create(
                    pk=entry["pk"],
                    defaults={
                        "client_id": entry["client_id"], "phone_number": entry["phone"],
                        "phone_type": entry["phone_type"], "is_primary": entry["is_primary"],
                        "notes": entry["notes"], "company_id": entry["company_id"],
                    },
                )
                ClientPhoneNumber.objects.filter(pk=entry["pk"]).update(
                    phone_number=entry["phone"], phone_normalized=canonical_phone_key(entry["phone"]) or ""
                )
        self.stdout.write(f"Restored {len(undo)} entries.")
